#!/usr/bin/env python3
"""Build the ISOM 472 team roster workbook.

Usage:  python3 build_roster.py

Regenerates isom472-team-roster.xlsx next to this script, then reopens the
saved file and verifies it (sheet names, dropdown contents, column widths,
and the instructor's email address).

Requires openpyxl.  No macros, no external links, no colour beyond grey
header shading, so the file opens cleanly in both Excel and Numbers.
"""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
OUT = HERE / "isom472-team-roster.xlsx"

TOKEN = "mo.almarzouq@ku.edu.kw"

STANDING_ROLES = [
    "Client Lead",
    "Design Lead",
    "Data Lead",
    "Build Lead",
    "Quality and FinOps Lead",
    "FinOps Lead",
    "Quality Lead",
]

# ---------------------------------------------------------------- styling ---

FONT = "Calibri"

TITLE = Font(name=FONT, size=14, bold=True)
SECTION = Font(name=FONT, size=11, bold=True)
HEAD = Font(name=FONT, size=11, bold=True)
BODY = Font(name=FONT, size=11)
NOTE = Font(name=FONT, size=10, italic=True, color="404040")
LABEL = Font(name=FONT, size=11, bold=True)
MONO_HINT = Font(name=FONT, size=11, bold=True)

HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
SECTION_FILL = PatternFill("solid", fgColor="F2F2F2")

thin = Side(style="thin", color="808080")
medium = Side(style="medium", color="404040")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
LEFT_MID = Alignment(horizontal="left", vertical="center", wrap_text=False)
LEFT_MID_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
CTR = Alignment(horizontal="center", vertical="center", wrap_text=False)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def put(ws, ref, value, font=BODY, align=LEFT_MID, fill=None, border=None):
    cell = ws[ref]
    cell.value = value
    cell.font = font
    cell.alignment = align
    if fill is not None:
        cell.fill = fill
    if border is not None:
        cell.border = border
    return cell


def merge_put(
    ws, ref, value, font=BODY, align=LEFT_MID, fill=None, border=None, height=None
):
    ws.merge_cells(ref)
    first = ref.split(":")[0]
    cell = put(ws, first, value, font, align, fill, border)
    if border is not None:
        for row in ws[ref]:
            for c in row:
                c.border = border
    if height is not None:
        ws.row_dimensions[ws[first].row].height = height
    return cell


def a4(ws, landscape=False):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = False
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.6


# ------------------------------------------------------------- sheet one ---


def build_roster(ws):
    widths(ws, {"A": 15, "B": 26, "C": 14, "D": 32, "E": 28, "F": 26})

    merge_put(ws, "A1:F1", "ISOM 472 — Team Roster", TITLE, LEFT_MID, height=22)
    merge_put(
        ws,
        "A2:F2",
        "Fall 2026. The Phase 1 lead fills this in and emails it to the "
        'instructor. See the "How to fill this in" sheet.',
        NOTE,
        LEFT_MID_WRAP,
        height=16,
    )

    merge_put(ws, "A4:B4", "Team name", LABEL, LEFT_MID)
    merge_put(ws, "C4:D4", None, BODY, LEFT_MID, border=BOX)
    merge_put(ws, "A5:B5", "Date submitted", LABEL, LEFT_MID)
    merge_put(ws, "C5:D5", None, BODY, LEFT_MID, border=BOX)
    put(ws, "E5", "Format: DD/MM/YYYY", NOTE, LEFT_MID)

    merge_put(
        ws, "A7:F7", "Team members", SECTION, LEFT_MID, fill=SECTION_FILL, height=18
    )
    merge_put(
        ws,
        "A8:F8",
        "One row per member. A team has four to six members. Rows 5 and 6 "
        "are optional — leave them empty if the team has four members.",
        NOTE,
        LEFT_MID_WRAP,
        height=16,
    )

    headers = [
        ("A9", "Member"),
        ("B9", "Full name"),
        ("C9", "Student ID"),
        ("D9", "KU email"),
        ("E9", "GitHub username\n(not the email address)"),
        ("F9", "Standing role"),
    ]
    for ref, text in headers:
        put(
            ws,
            ref,
            text,
            HEAD,
            Alignment(horizontal="left", vertical="center", wrap_text=True),
            fill=HEAD_FILL,
            border=BOX,
        )
    ws["A9"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[9].height = 32

    for i in range(6):
        row = 10 + i
        ws.row_dimensions[row].height = 20
        optional = i >= 4
        # Rows 5 and 6 are marked optional in the member column itself, so the
        # marking prints. A gutter column outside the print area would not.
        label = f"{i + 1} — optional" if optional else f"{i + 1}"
        put(ws, f"A{row}", label, NOTE if optional else BODY, LEFT_MID, border=BOX)
        for col in "BCDEF":
            put(ws, f"{col}{row}", None, BODY, LEFT_MID, border=BOX)

    # Make the GitHub username column visually prominent: bold entry cells and
    # a medium box drawn around the whole column, header included.
    for row in range(10, 16):
        ws[f"E{row}"].font = MONO_HINT
    for row in range(9, 16):
        cell = ws[f"E{row}"]
        left = medium
        right = medium
        top = medium if row == 9 else thin
        bottom = medium if row == 15 else thin
        cell.border = Border(left=left, right=right, top=top, bottom=bottom)

    merge_put(
        ws,
        "A16:F16",
        "GitHub username, not the email address. Example username: "
        "s-almutairi  —  example email: s.almutairi@ku.edu.kw",
        NOTE,
        LEFT_MID_WRAP,
        height=16,
    )

    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(STANDING_ROLES) + '"',
        allow_blank=True,
        showDropDown=False,  # False here means: DO show the in-cell dropdown
        showErrorMessage=True,
        showInputMessage=True,
    )
    dv.promptTitle = "Standing role"
    dv.prompt = "Pick one role from the list. See the Roles sheet."
    dv.errorTitle = "Pick a role from the list"
    dv.error = "Choose one of the roles offered in the dropdown."
    ws.add_data_validation(dv)
    dv.add("F10:F15")

    merge_put(
        ws, "A18:F18", "Phase leads", SECTION, LEFT_MID, fill=SECTION_FILL, height=18
    )
    merge_put(
        ws,
        "A19:F19",
        "One member leads each phase, on top of their standing role. A "
        "team of six has everyone lead once; a smaller team repeats "
        "members. Type the name — there is no list to pick from.",
        NOTE,
        LEFT_MID_WRAP,
        height=28,
    )

    merge_put(ws, "A20:B20", "Phase", HEAD, LEFT_MID, fill=HEAD_FILL, border=BOX)
    merge_put(ws, "C20:D20", "Member name", HEAD, LEFT_MID, fill=HEAD_FILL, border=BOX)
    for i in range(6):
        row = 21 + i
        ws.row_dimensions[row].height = 20
        merge_put(ws, f"A{row}:B{row}", f"Phase {i + 1}", BODY, LEFT_MID, border=BOX)
        merge_put(ws, f"C{row}:D{row}", None, BODY, LEFT_MID, border=BOX)

    ws.freeze_panes = "A10"
    a4(ws, landscape=True)
    ws.print_area = "A1:F26"


# ------------------------------------------------------------- sheet two ---

ROLE_ROWS = [
    (
        "Client Lead",
        "The client relationship, and what the system must do.",
        "The backlog: user stories with acceptance criteria.",
    ),
    (
        "Design Lead",
        "What it looks like and how it flows.",
        "The HTML prototype and the screen list.",
    ),
    (
        "Data Lead",
        "The data and its structure.",
        "The schema and seed data in Supabase.",
    ),
    (
        "Build Lead",
        "That it actually runs.",
        "The deployed system at its public address, plus release notes each sprint.",
    ),
    (
        "FinOps Lead",
        "What the AI work costs, and what the team changed as a result.",
        "The Sprint FinOps Ledger.",
    ),
    (
        "Quality Lead",
        "Finding defects in the running application.",
        "Bug issues filed in the repository with steps to reproduce, tracked to "
        "a disposition.",
    ),
    (
        "Phase Lead\n(rotates: one per phase, held on top of a standing role)",
        "Running the phase, and being the instructor's point of contact for it.",
        "The phase delivery note.",
    ),
]

SIZING = [
    "A team of four takes the first four roles: Client Lead, Design Lead, "
    "Data Lead, Build Lead.",
    "A fifth member takes the combined Quality and FinOps Lead.",
    "At six members that combined role splits into two: Quality Lead and FinOps Lead.",
    "In a team of four, the Phase Lead also keeps that phase's ledger and files that phase's bug tickets, so both rotate with the lead.",
]

EVERYONE = [
    "Every member delivers stories end to end.",
    "Every member reviews teammates' pull requests. Nobody approves their own.",
]


def build_roles(ws):
    widths(ws, {"A": 34, "B": 46, "C": 52})

    merge_put(ws, "A1:C1", "Roles", TITLE, LEFT_MID, height=22)
    merge_put(
        ws,
        "A2:C2",
        "Every member holds exactly one standing role for the whole "
        "semester. The Phase Lead rotates and sits on top of it.",
        NOTE,
        LEFT_MID_WRAP,
        height=16,
    )

    for ref, text in (
        ("A4", "Role"),
        ("B4", "Owns — the one thing"),
        ("C4", "Artifact it produces"),
    ):
        put(ws, ref, text, HEAD, LEFT_MID_WRAP, fill=HEAD_FILL, border=BOX)
    ws.row_dimensions[4].height = 18

    row = 5
    for role, owns, artifact in ROLE_ROWS:
        put(ws, f"A{row}", role, HEAD, LEFT_TOP, border=BOX)
        put(ws, f"B{row}", owns, BODY, LEFT_TOP, border=BOX)
        put(ws, f"C{row}", artifact, BODY, LEFT_TOP, border=BOX)
        longest = max(
            len(owns) / 44.0,
            len(artifact) / 50.0,
            len(role) / 32.0,
            role.count("\n") + 1,
        )
        ws.row_dimensions[row].height = max(30, 15 * (int(longest) + 1))
        row += 1

    row += 1
    merge_put(
        ws,
        f"A{row}:C{row}",
        "How many members, how many roles",
        SECTION,
        LEFT_MID,
        fill=SECTION_FILL,
        height=18,
    )
    row += 1
    for line in SIZING:
        merge_put(
            ws,
            f"A{row}:C{row}",
            "•  " + line,
            BODY,
            LEFT_TOP,
            height=17 if len(line) < 110 else 32,
        )
        row += 1

    row += 1
    merge_put(
        ws,
        f"A{row}:C{row}",
        "Holds for everyone, whatever the role",
        SECTION,
        LEFT_MID,
        fill=SECTION_FILL,
        height=18,
    )
    row += 1
    for line in EVERYONE:
        merge_put(ws, f"A{row}:C{row}", "•  " + line, BODY, LEFT_TOP, height=17)
        row += 1

    ws.freeze_panes = "A5"
    a4(ws, landscape=True)
    ws.print_area = f"A1:C{row}"


# ----------------------------------------------------------- sheet three ---

STEPS = [
    "Form your team. Four to six members.",
    'Agree the standing roles. Read the "Roles" sheet first. Every member '
    "holds exactly one standing role.",
    "Agree who leads each of the six phases. One member per phase, on top of "
    "their standing role.",
    "Every member creates a GitHub account at github.com and gives the Phase 1 "
    "lead their GitHub username.",
    'Fill in the "Team roster" sheet: team name, date, one row per member, '
    "and the six phase leads. Pick each standing role from the dropdown in the "
    "last column.",
    "Save the file. Rename it to include your team name, for example "
    "isom472-team-roster-northgate.xlsx",
    "Email the saved file to " + TOKEN + " as an attachment.",
]


def build_how(ws):
    widths(ws, {"A": 5, "B": 96})

    merge_put(ws, "A1:B1", "How to fill this in", TITLE, LEFT_MID, height=22)
    merge_put(
        ws,
        "A2:B2",
        "The Phase 1 lead does this once, for the whole team.",
        NOTE,
        LEFT_MID,
        height=16,
    )

    row = 4
    for i, step in enumerate(STEPS, start=1):
        put(ws, f"A{row}", f"{i}.", HEAD, Alignment(horizontal="right", vertical="top"))
        put(ws, f"B{row}", step, BODY, LEFT_TOP)
        ws.row_dimensions[row].height = 16 * (len(step) // 92 + 1)
        row += 1

    row += 1
    merge_put(
        ws,
        f"A{row}:B{row}",
        "Deadline",
        SECTION,
        LEFT_MID,
        fill=SECTION_FILL,
        height=18,
    )
    row += 1
    merge_put(
        ws,
        f"A{row}:B{row}",
        "Before you leave the setup session. Exact dates are announced in "
        "class and on the course site.",
        BODY,
        LEFT_TOP,
        height=17,
    )
    row += 1
    merge_put(
        ws,
        f"A{row}:B{row}",
        "The instructor creates your team's GitHub repository once this "
        "file arrives. Until it arrives, the team has no repository.",
        BODY,
        LEFT_TOP,
        height=17,
    )

    row += 2
    merge_put(
        ws,
        f"A{row}:B{row}",
        "The one field teams get wrong",
        SECTION,
        LEFT_MID,
        fill=SECTION_FILL,
        height=18,
    )
    row += 1
    merge_put(
        ws,
        f"A{row}:B{row}",
        "A GitHub username is not an email address. The username is the "
        "short name you chose when you signed up, and it is what appears "
        "in the address of your GitHub page.",
        BODY,
        LEFT_TOP,
        height=17,
    )
    row += 1
    merge_put(
        ws,
        f"A{row}:B{row}",
        "GitHub username:  s-almutairi        KU email:  s.almutairi@ku.edu.kw",
        HEAD,
        LEFT_TOP,
        height=17,
    )
    row += 1
    merge_put(
        ws,
        f"A{row}:B{row}",
        "The roster needs the username. Put the email in the KU email "
        "column and the username in the GitHub username column.",
        BODY,
        LEFT_TOP,
        height=17,
    )

    ws.freeze_panes = "A4"
    a4(ws, landscape=False)
    ws.print_area = f"A1:B{row}"


# ------------------------------------------------------------------ build ---


def build():
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Team roster"
    build_roster(ws1)
    build_roles(wb.create_sheet("Roles"))
    build_how(wb.create_sheet("How to fill this in"))
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb.active = 0
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    return OUT


# ----------------------------------------------------------------- verify ---


def effective_width(ws, cell):
    """Column width available to a cell, summed across any merge it is in."""
    for rng in ws.merged_cells.ranges:
        if (
            rng.min_row <= cell.row <= rng.max_row
            and rng.min_col <= cell.column <= rng.max_col
        ):
            return sum(
                ws.column_dimensions[get_column_letter(c)].width or 8.43
                for c in range(rng.min_col, rng.max_col + 1)
            )
    return ws.column_dimensions[get_column_letter(cell.column)].width or 8.43


def verify(path):
    problems = []
    wb = load_workbook(path)
    names = wb.sheetnames
    expected = ["Team roster", "Roles", "How to fill this in"]
    if names != expected:
        problems.append(f"sheet names {names} != {expected}")

    # dropdown
    dvs = wb["Team roster"].data_validations.dataValidation
    if len(dvs) != 1:
        problems.append(f"expected 1 data validation, found {len(dvs)}")
    else:
        dv = dvs[0]
        offered = dv.formula1.strip('"').split(",")
        if offered != STANDING_ROLES:
            problems.append(f"dropdown offers {offered}")
        dv_range = str(dv.sqref)

    # overflow: an unwrapped cell whose text is wider than its column
    overflow = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                text = str(cell.value)
                if cell.alignment and cell.alignment.wrap_text:
                    continue
                longest = max(len(seg) for seg in text.split("\n"))
                avail = effective_width(ws, cell)
                bold = cell.font and cell.font.bold
                need = longest * (1.06 if bold else 1.0)
                if need > avail:
                    overflow.append(
                        f"{ws.title}!{cell.coordinate} needs ~{need:.0f} "
                        f"has {avail:.0f}: {text[:40]!r}"
                    )
    if overflow:
        problems.extend(overflow)

    # nothing populated may sit outside the print area, or it will not print
    import re

    from openpyxl.utils import column_index_from_string

    for ws in wb.worksheets:
        pa = ws.print_area
        if isinstance(pa, (list, tuple)):
            pa = pa[0] if pa else None
        m = re.search(r"\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)", pa or "")
        if not m:
            problems.append(f"{ws.title}: no usable print area ({pa!r})")
            continue
        max_col = column_index_from_string(m.group(3))
        max_row = int(m.group(4))
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if cell.column > max_col or cell.row > max_row:
                    problems.append(
                        f"{ws.title}!{cell.coordinate} is outside the print area {pa}"
                    )

    # token count
    hits = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    hits += str(cell.value).count(TOKEN)
    if hits != 1:
        problems.append(f"token {TOKEN} appears {hits} times, expected 1")

    dims = {ws.title: ws.dimensions for ws in wb.worksheets}
    return problems, names, dims, dv_range, offered


if __name__ == "__main__":
    p = build()
    problems, names, dims, dv_range, offered = verify(p)
    print(f"wrote {p}  ({p.stat().st_size} bytes)")
    for n in names:
        print(f"  sheet {n!r}: {dims[n]}")
    print(f"  dropdown on {dv_range}: {offered}")
    if problems:
        print("PROBLEMS:")
        for q in problems:
            print("  -", q)
    else:
        print("VERIFIED: no problems")
